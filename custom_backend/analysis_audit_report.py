"""Portable method/parameter/result audit appendices for server-computed analyses."""
from __future__ import annotations

import csv
import json
import re
import textwrap
from io import BytesIO, StringIO

from fastapi import Response


def _fields(value, path=''):
    if isinstance(value, dict) and value:
        for key, item in value.items():
            yield from _fields(item, f'{path}.{key}' if path else key)
    elif isinstance(value, list) and value:
        for index, item in enumerate(value):
            yield from _fields(item, f'{path}[{index}]')
    else:
        yield path, json.dumps(value, ensure_ascii=False, allow_nan=False) if not isinstance(value, str) else value


def _latex_escape(value) -> str:
    escaped = {'\\': r'\textbackslash{}', '{': r'\{', '}': r'\}', '$': r'\$',
               '&': r'\&', '%': r'\%', '#': r'\#', '_': r'\_',
               '~': r'\textasciitilde{}', '^': r'\textasciicircum{}'}
    return ''.join(escaped.get(char, char) for char in str(value)).replace('\n', r'\newline{}')


def _latex_breakable(value, min_length=24, run_limit=6) -> str:
    parts = []
    for token in re.split(r'(\s+)', str(value)):
        if len(token) < min_length:
            parts.append(_latex_escape(token))
            continue
        run = 0
        for char in token:
            run = run + 1 if char.isascii() and char.isalnum() else 0
            parts.append(_latex_escape(char))
            if char in '._/:-[]' or (run_limit and run >= run_limit):
                parts.append(r'\allowbreak{}')
                run = 0
    return ''.join(parts)


def audit_response(result: dict, format: str, title: str, tables=(), filename='reviewflow_analysis_audit') -> Response:
    """Keep exact numeric values and every provenance field; never accept participant rows."""
    tables = tuple(tables)
    fields = list(_fields(result))
    table_fields = [field for index, summary in enumerate(tables)
                    for field in _fields(summary, f'tables[{index}]')]
    if any(part in {'participants', 'participant_rows'} for path, _ in [*fields, *table_fields]
           for part in path.replace('[', '.').split('.')):
        raise ValueError('participant rows must not be included in audit reports')
    reminder = 'Verify the original sources and cite the methods used. Field paths preserve the calculation parameters and source snapshots.'
    if format == 'tex':
        lines = [r'% Compile with XeLaTeX (ctex package required).',
                 r'\documentclass[UTF8]{ctexart}',
                 r'\usepackage[margin=24mm]{geometry}',
                 r'\usepackage{longtable,array}',
                 r'\setlength{\emergencystretch}{3em}',
                 r'\begin{document}',
                 r'\section*{' + _latex_escape(title) + '}',
                 _latex_escape(reminder) + r'\par']
        for summary in tables:
            lines.extend([r'\subsection*{' + _latex_escape(summary['title']) + '}',
                          r'\begin{longtable}{@{}' + r'@{\hspace{.01\linewidth}}'.join(r'>{\raggedright\arraybackslash}p{' + str(round(.90/len(summary['columns']), 3)) + r'\linewidth}' for _ in summary['columns']) + '@{}}',
                          ' & '.join(map(_latex_escape, summary['columns'])) + r' \\ \hline'])
            lines.extend(' & '.join(_latex_breakable(value, 8) for value in row) + r' \\ \hline'
                         for row in summary['rows'])
            lines.extend([r'\end{longtable}', _latex_escape(summary['note']) + r'\par'])
        lines.extend([r'\subsection*{Full method, parameter and source audit}',
                      r'\begin{longtable}{@{}>{\raggedright\arraybackslash}p{.31\linewidth}@{}>{\raggedright\arraybackslash}p{.63\linewidth}@{}}',
                      r'Field & Value / method / source \\ \hline'])
        lines.extend(_latex_breakable(path, 0, 0) + ' & ' + _latex_breakable(value, 40) + r' \\ \hline'
                     for path, value in fields)
        lines.extend([r'\end{longtable}', r'\end{document}'])
        content = ('\n'.join(lines) + '\n').encode('utf-8')
        media = 'application/x-tex; charset=utf-8'
    elif format == 'csv':
        stream = StringIO(newline='')
        writer = csv.writer(stream)
        writer.writerow(['Report', 'Field', 'Value'])
        # Spreadsheet formula injection applies to labels and source text as well as values.
        def cell(value):
            if value.lstrip().startswith('-'):
                try:
                    if isinstance(json.loads(value), (int, float)):
                        return value
                except ValueError:
                    pass
            return "'" + value if value.lstrip().startswith(('=', '+', '-', '@')) else value
        for path, value in [('report_note', reminder), *table_fields, *fields]:
            writer.writerow([cell(title), cell(path), cell(value)])
        content = ('\ufeff' + stream.getvalue()).encode('utf-8')
        media = 'text/csv; charset=utf-8'
    elif format == 'docx':
        from docx import Document
        doc = Document()
        doc.add_heading(title, 0)
        doc.add_paragraph(reminder)
        for summary in tables:
            doc.add_heading(summary['title'], 1)
            table = doc.add_table(rows=1, cols=len(summary['columns']))
            table.style = 'Table Grid'
            for cell, label in zip(table.rows[0].cells, summary['columns']):
                cell.text = label
            for row in summary['rows']:
                for cell, value in zip(table.add_row().cells, row):
                    cell.text = str(value)
            doc.add_paragraph(summary['note'])
        doc.add_heading('Full method, parameter and source audit', 1)
        section = None
        table = None
        for path, value in fields:
            group = path.split('.')[0].split('[')[0]
            if group != section:
                section = group
                doc.add_heading(group.replace('_', ' '), 1)
                table = doc.add_table(rows=1, cols=2)
                table.style = 'Table Grid'
                table.rows[0].cells[0].text = 'Field'
                table.rows[0].cells[1].text = 'Value / method / source'
            cells = table.add_row().cells
            cells[0].text, cells[1].text = path, value
        stream = BytesIO()
        doc.save(stream)
        content = stream.getvalue()
        media = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    elif format == 'pptx':
        from pptx import Presentation
        from pptx.util import Inches, Pt
        presentation = Presentation()
        presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
        for summary in tables:
            for offset in range(0, len(summary['rows']), 8):
                slide = presentation.slides.add_slide(presentation.slide_layouts[6])
                heading = slide.shapes.add_textbox(Inches(.5), Inches(.25), Inches(12.3), Inches(.7))
                heading.text_frame.paragraphs[0].text = summary['title'] + f' · {offset//8+1}'
                heading.text_frame.paragraphs[0].font.size = Pt(24)
                rows = summary['rows'][offset:offset+8]
                table = slide.shapes.add_table(len(rows)+1,len(summary['columns']),Inches(.5),Inches(1.1),Inches(12.3),Inches(.5*(len(rows)+1))).table
                for i, row in enumerate([summary['columns'], *rows]):
                    for j, value in enumerate(row):
                        cell = table.cell(i,j)
                        cell.text = str(value)
                        for paragraph in cell.text_frame.paragraphs:
                            paragraph.font.size = Pt(14)
                footnote = slide.shapes.add_textbox(Inches(.5), Inches(5.9), Inches(12.3), Inches(1.3))
                footnote.text_frame.word_wrap = True
                footnote.text_frame.paragraphs[0].text = summary['note']
                footnote.text_frame.paragraphs[0].font.size = Pt(12)
        lines = [reminder, '']
        for path, value in fields:
            lines.extend(textwrap.wrap(f'{path}: {value}', width=48, replace_whitespace=False) or [''])
        for offset in range(0, len(lines), 16):
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            heading = slide.shapes.add_textbox(Inches(.5), Inches(.25), Inches(12.3), Inches(.6))
            heading.text_frame.paragraphs[0].text = f'{title} · {offset//16+1}'
            heading.text_frame.paragraphs[0].font.size = Pt(24)
            body = slide.shapes.add_textbox(Inches(.5), Inches(1), Inches(12.3), Inches(5.9))
            body.text_frame.word_wrap = True
            for index, line in enumerate(lines[offset:offset+16]):
                paragraph = body.text_frame.paragraphs[0] if index == 0 else body.text_frame.add_paragraph()
                paragraph.text = line
                paragraph.font.size = Pt(16)
                paragraph.space_after = Pt(4)
            footer = slide.shapes.add_textbox(Inches(.5), Inches(7), Inches(12.3), Inches(.3))
            footer.text_frame.paragraphs[0].text = 'Method, parameter and source audit — retain all pages with the reported results.'
            footer.text_frame.paragraphs[0].font.size = Pt(10)
        stream = BytesIO()
        presentation.save(stream)
        content = stream.getvalue()
        media = 'application/vnd.openxmlformats-officedocument.presentationml.presentation'
    else:
        raise ValueError('unsupported audit format')
    return Response(content=content, media_type=media,
                    headers={'Content-Disposition': f'attachment; filename={filename}.{format}'})


def sensitivity_differences(result: dict, view: dict) -> dict:
    """Presentation arithmetic only; compare cached intervals with the same target."""
    cells = [*result['cells'], *[
        {**cell, 'key':cell['model']+'_bootstrap_'+cell['interval_type'],
         'applicable':True, 'interval_method':'bootstrap_'+cell['interval_type']}
        for cell in result.get('bootstrap') or []]]
    if result.get('pooled_effect_profile', {}).get('applicable'):
        cells.append(result['pooled_effect_profile'])
    cells = [cell for cell in cells if cell['applicable'] and
             (cell['interval_method'] == 'prediction') == (view['target'] == 'prediction')]
    base = next((cell for cell in cells if cell['key'] == view['reference']), None)
    if base is None:
        raise ValueError('comparison reference must name an available interval with the selected target')
    threshold = view['width_factor_threshold']
    records = []
    for cell in cells:
        ratio = cell['width']/base['width'] if base['width'] > 0 else None
        factor = max(ratio,1/ratio) if ratio is not None and ratio > 0 else None
        records.append(dict(key=cell['key'],estimate_difference=cell['estimate']-base['estimate'],
                            width_ratio=ratio,width_factor=factor,
                            exceeds_threshold=factor is not None and threshold is not None and factor > threshold))
    return dict(settings=view,analysis_scale=result['analysis_scale'],rows=records,
                note='Width factor = max(width/reference, reference/width). User-selected descriptive threshold; not a significance test or model-selection recommendation. Consult and cite Cochrane Handbook section 10.14.')
