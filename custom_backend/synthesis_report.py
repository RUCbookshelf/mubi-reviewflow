"""DOCX and PPTX exports for a server-computed synthesis result."""

from __future__ import annotations

import json
import math
import textwrap
from io import BytesIO

_LOG_SCALE_MEASURES = {"RR", "OR", "HR", "RATE_RATIO"}
_POINT_ESTIMATES = {"MEAN", "LOGIT_PROP", "LOG_RATE"}


def _number(value: float | int | None) -> str:
    return "—" if value is None else f"{float(value):.4g}"


def _analysis_estimate(effect: dict) -> float:
    estimate = float(effect["estimate"])
    return math.log(estimate) if effect["measure"] in _LOG_SCALE_MEASURES else estimate


def _display_value(value: float, measure: str) -> float:
    if measure in _LOG_SCALE_MEASURES or measure in {"LOG_RATE", "LOG_ROM", "PAIRED_OR"}:
        return math.exp(value)
    if measure == "FISHER_Z":
        return math.tanh(value)
    if measure == "LOGIT_PROP":
        return 1 / (1 + math.exp(-value)) if value >= 0 else math.exp(value) / (1 + math.exp(value))
    return value


def _bounds(report: dict) -> tuple[float, float]:
    result = report["result"]
    values = [result["ci_analysis_low"], result["ci_analysis_high"]]
    if result["measure"] not in _POINT_ESTIMATES:
        values.append(0.0)
    for effect in result["effects"]:
        center = _analysis_estimate(effect)
        values.extend((center - 1.96 * effect["se"], center + 1.96 * effect["se"]))
    low, high = min(values), max(values)
    padding = (high - low) * 0.08 or max(abs(low), 1.0) * 0.08
    return low - padding, high + padding


def _plot_bar(low: float, estimate: float, high: float, bounds: tuple[float, float],
              width: int = 35, reference: bool = True) -> str:
    minimum, maximum = bounds

    def pos(value: float) -> int:
        return max(0, min(width - 1, round((value - minimum) / (maximum - minimum) * (width - 1))))

    start, center, end, reference_pos = pos(low), pos(estimate), pos(high), pos(0)
    line = [" "] * width
    for index in range(start, end + 1):
        line[index] = "-"
    if reference:
        line[reference_pos] = "|"
    line[center] = "*"
    return "".join(line)


def _citations(report: dict) -> list[str]:
    result = report["result"]
    reminders = [
        "Core inverse-variance synthesis: Harrer et al. (2022), Doing Meta-Analysis with R, chapters 3–5; Borenstein et al. (2021), Introduction to Meta-Analysis, 2nd ed., chapters 11–14.",
        "Descriptive I²: Higgins & Thompson (2002), Statistics in Medicine 21:1539–1558; derive it from fixed inverse-variance Q, irrespective of the selected pooling model. Verify the original method and cite the τ² estimator used.",
        "Cite the selected study reports and verify effect direction, outcome, time point, units, and saved input values against each source.",
    ]
    if result["measure"] == "MEAN":
        reminders.append("One-group mean: Harrer et al. (2022), Doing Meta-Analysis with R, section 3.2.1; check the original text and outcome unit.")
    elif result["measure"] == "LOGIT_PROP":
        reminders.append("One-group logit proportion: Harrer et al. (2022), sections 3.2.2 and 4.2.6; this inverse-variance method needs non-boundary study counts.")
    elif result["measure"] == "LOG_RATE":
        reminders.append("One-group log incidence rate: metafor escalc documentation, measure IRLN (https://wviechtb.github.io/metafor/reference/escalc.html); verify the common person-time unit.")
    elif result["measure"] in {"SMCR", "SMCC"}:
        reminders.append("One-group standardized post-minus-pre change: Borenstein et al. (2021), ch. 4, pp. 28–29; metafor escalc SMCR/SMCRH/SMCC. Check each saved standardizer and variance_method; cite Becker (1988), Bonett (2008), or Gibbons et al. (1993) as used. The SMCR variance follows metafor rather than book equation 4.28. This is not a controlled treatment effect.")
    glass_effects = [effect for effect in result["effects"] if effect.get("entry_method") == "glass_delta_two_arm" or
                     effect.get("input_data", {}).get("standardizer") == "control_arm_sd"]
    if any(effect.get("entry_method") == "glass_delta_two_arm" for effect in glass_effects):
        reminders.append("Glass delta: control-arm SD standardizer, no small-sample correction, large-sample variance (metafor SMD1, correct=FALSE, vtype=LS). Verify and cite Harrer et al. (2022), section 4.2.2, pp. 112–114, and Hedges (1982). These effects are kept separate from Hedges g.")
    if any(effect.get("entry_method") == "manual" for effect in glass_effects):
        reminders.append("Reported Glass delta: control-arm SD standardizer with a source-reported SE. Verify the report's bias correction and variance method; do not infer that it used metafor SMD1. Check Harrer et al. (2022), section 4.2.2, pp. 112–114, and cite the actual method and study reports.")
    if result["measure"] == "SMD" and not glass_effects:
        if any(effect.get("entry_method") != "manual" for effect in result["effects"]):
            reminders.append("ReviewFlow-calculated independent-arm SMD uses a pooled within-group SD and Hedges g small-sample correction: Hedges (1981), DOI 10.3102/10769986006002107. Verify the study inputs and cite the method and reports.")
        if any(effect.get("entry_method") == "manual" for effect in result["effects"]):
            reminders.append("Reported Hedges g: pooled within-group SD with a source-reported SE. Verify the report's small-sample correction, effect direction and variance method; cite Hedges (1981), DOI 10.3102/10769986006002107, when confirmed, along with the study reports.")
    survival_summaries = [effect for effect in result["effects"] if effect.get("input_data", {}).get("survival_summary_conversion")]
    if survival_summaries:
        reminders.append("Median-survival conversion assumes constant exponential hazards and uses approximate Poisson log-scale uncertainty. Inspect event_sources for reported events versus explicit half-n approximations; censoring patterns are not modeled. Prefer reported HR/CI or log-rank data when available. Reconstructed O−E/V is not reported log-rank evidence.")
        reminders.extend(dict.fromkeys(source for effect in survival_summaries for source in effect["input_data"].get("method_sources", [])))
    estimated = [effect for effect in result["effects"] if effect.get("input_data", {}).get("estimated_summary")]
    if estimated:
        reminders.append("Formula-estimated arm means/SDs are not observed data. Standard effect SEs omit reconstruction uncertainty. Compare synthesis with and without these studies and across estimation methods; verify distribution assumptions and cite the original study reports.")
        for effect in estimated:
            reminders.append(f"Estimated summary: study {effect['study_id']}; result {effect.get('result_id')}; source {effect.get('source_key')}. Check saved arm_estimation for arm-specific formulas and original quantiles.")
        reminders.extend(dict.fromkeys(source for effect in estimated for source in effect["input_data"].get("method_sources", [])))
    if result["measure"] in {"PAIRED_OR", "PAIRED_RD"}:
        reminders.append("Paired binary effects: matched-pair OR uses log-scale estimates and SE, with exp-transformed display; paired RD uses original counts and natural-scale SE. Normal Wald intervals; inspect saved or_correction_applied and input counts for sensitivity corrections. Verify and cite Curtin, Elbourne & Altman (2002), Binary outcomes, DOI 10.1002/sim.1206, and Fagerland et al. (2014), DOI 10.1002/sim.6148 (interval limitations). Analyze separately from independent-arm effects.")
    if result["measure"] in {"PHI", "R_EQUIV_APPROX"}:
        reminders.append("Binary correlation on the natural scale: verify Bonett (2021), Volume 3, section 3.4, equations 3.11–3.12 and examples 3.6/3.8; cite the actual method and source reports. PHI uses the marginal-dependent large-sample SE, not the continuous-Pearson 1/sqrt(n-3) SE. R_EQUIV_APPROX is the Bonett & Price (2005) tetrachoric approximation, assuming latent bivariate-normal variables and adding 0.5 to every cell. These two measures and continuous-outcome Fisher z must be analyzed separately. Normal intervals can exceed correlation bounds; do not silently truncate them.")
    if result["measure"] == "LOG_ROM":
        reminders.append("Ratio of means for independent continuous-outcome arms: Borenstein et al. (2021), ch. 4, pp. 30–31; Hedges, Gurevitch & Curtis (1999), equation 1. Uncorrected log ROM and separate-arm delta variance; saved estimates and SE are log-scale, displayed ratios are exp-transformed. Do not interpret as binary risk ratios. Verify and cite the original methods.")
    if any(effect.get("entry_method") == "multi-arm combination" for effect in result["effects"]):
        reminders.append("Multi-arm combination: Harrer et al. (2022), Doing Meta-Analysis with R, sections 3.5.2 and 17.9, pp. 88–89 and 434–435; Cochrane Handbook, section 23.3.4. Verify arm eligibility and cite the study report.")
    if any(effect.get("entry_method") == "zero_cell_binary" for effect in result["effects"]):
        reminders.append('Zero-cell OR/RR: check the saved correction and adjusted 2×2 counts. A selected 0.5 correction adds 0.5 to all four cells only when a cell is zero (metafor escalc, add=0.5, to="only0": https://wviechtb.github.io/metafor/reference/escalc.html). Verify Borenstein et al. (2021), Introduction to Meta-Analysis, 2nd ed., ch. 5, pp. 34–37, and Cochrane Handbook ch. 10, §§10.4.4.1–10.4.4.2; cite the method and study reports.')
    methods = {effect.get("entry_method") or "" for effect in result["effects"]}
    if "correlation_r_n" in methods:
        reminders.append("Pearson r to Fisher z: z=atanh(r), SE(z)=1/sqrt(n−3); verify the reported correlation and sample size. Harrer et al. (2022), Doing Meta-Analysis with R, §3.2.3, pp. 62–63; metafor escalc ZCOR (https://wviechtb.github.io/metafor/reference/escalc.html). Cite the method and study report.")
    if "generic_wald_p" in methods:
        reminders.append("Exact two-sided normal Wald p to SE: SE=abs(effect on analysis scale)/norm.isf(p/2). Ratio effects use their log scale. This does not apply to t, score, likelihood-ratio or threshold p values. Verify Cochrane Handbook ch. 6, §§6.3.1–6.3.2 (https://www.cochrane.org/authors/handbooks-and-manuals/handbook/current/chapter-06), and cite the method and study report.")
    if methods.intersection({"generic_ci", "hazard_ratio_ci"}):
        reminders.append("Reported CI to SE: divide the interval width on the analysis scale by twice the normal critical value at its reported confidence level; ratios use log limits. This assumes a symmetric normal/Wald interval on that scale. Verify Cochrane Handbook ch. 6, §§6.3.1–6.3.2 (https://www.cochrane.org/authors/handbooks-and-manuals/handbook/current/chapter-06), and cite the method and study report.")
    direct = [effect for effect in result["effects"] if effect.get("entry_method") in
              {"single_mean_direct", "single_proportion_direct", "single_rate_direct"}]
    if any(effect.get("input_data", {}).get("ci_method") == "wald_normal" for effect in direct):
        reminders.append("One-group reported normal/Wald CI to SE: transform proportion/rate limits to the stated analysis scale, then divide their width by twice the normal critical value at the saved confidence level. Exact, score, profile and t intervals are unsuitable. Check Cochrane Handbook ch. 6, §§6.3.1–6.3.2 (https://www.cochrane.org/authors/handbooks-and-manuals/handbook/current/chapter-06); verify the original interval method and cite the study report.")
    if any(effect["measure"] in {"LOGIT_PROP", "LOG_RATE"} and
           effect.get("input_data", {}).get("precision_scale") == "raw" and
           ("se" in effect.get("input_data", {}) or "variance" in effect.get("input_data", {}))
           for effect in direct):
        reminders.append("One-group raw SE or variance was converted to the logit/log scale by a first-order delta approximation. Check Harrer et al. (2022), sections 3.2.2 and 3.3.3, or metafor escalc documentation; verify the original uncertainty scale and cite the method and report.")
    if "rate_ratio_events_time" in methods:
        reminders.append("Incidence rate ratio from two event counts and person-times: log-scale SE=sqrt(1/events_t+1/events_c), requiring positive counts in both arms. Verify Harrer et al. (2022), Doing Meta-Analysis with R, §3.3.3, pp. 76–79, and cite the method and study report.")
    if "survival_logrank_oe_v" in methods:
        reminders.append("Log-rank O−E/V to treatment/control HR: Tierney et al. (2007), DOI 10.1186/1745-6215-8-16, equation 7; Cochrane Handbook ch. 6, §6.8.2. Verify the reported arm direction, O−E and V, and cite the study report.")
    if "crossover_2x2_md" in methods:
        reminders.append("Two-sequence crossover MD: verify complete-pair sequence contrasts, period and carryover assumptions against Cochrane Handbook ch. 23, §§23.2.3 and 23.2.5, and Cooper, Hedges & Valentine (2019), Handbook of Research Synthesis and Meta-Analysis, 3rd ed., §§11.2 and 21.3.3.7; cite the study report.")
    if any(method.startswith("paired_md_") for method in methods):
        reminders.append("Paired mean difference: verify the complete-pair difference SD or within-pair correlation. Borenstein et al. (2021), Introduction to Meta-Analysis, 2nd ed., ch. 4, pp. 23–29; Cooper, Hedges & Valentine (2019), ch. 11, §§11.2.1.2–11.2.2.2. Cite the method and study report.")
    if any(method.startswith("parallel_change_") for method in methods):
        reminders.append("Parallel-arm change-score MD: use each arm's analyzed paired n and change SD or documented baseline–post correlation; this is a controlled difference in changes. Borenstein et al. (2021), ch. 4, pp. 23–29; Cochrane Handbook ch. 6, §6.5.2.8 and ch. 10, §10.5.2. Cite the method and study report.")
    if methods.intersection({"independent_t_n", "independent_p_n", "independent_f_n", "independent_d_n"}):
        reminders.append('Independent-group test statistic to Hedges g: Borenstein et al. (2021), Introduction to Meta-Analysis, 2nd ed., ch. 4, pp. 26–28, equations 4.20–4.25. The sampling variance follows metafor escalc SMD vtype="LS2" (https://wviechtb.github.io/metafor/reference/escalc.html), not its default LS variance. Verify pooled-variance design and the reported sign; cite the method and study report.')
    cluster_effects = [effect for effect in result["effects"] if effect.get("entry_method") in
                       {"cluster_trial_design_effect", "cluster_trial_variable_size_design_effect"}]
    if cluster_effects:
        reminders.append("Cluster-trial design effects are approximate: verify the original arm summaries, ICC and its source, cluster sizes, and the trial's own analysis before citing the method.")
        reminders.extend(dict.fromkeys(source for effect in cluster_effects
                                       for source in effect.get("input_data", {}).get("method_sources", [])))
    if result["model"] == "random":
        reminders.append("Random-effects estimator: DerSimonian and Laird (1986), DOI 10.1016/0197-2456(86)90046-2.")
    elif result["model"] == "random_pm":
        reminders.append("Random-effects estimator: Paule and Mandel (1982), Journal of Research of the National Bureau of Standards, 87, 377–385.")
    elif result["model"] == "random_reml":
        reminders.append("Restricted maximum likelihood (REML) heterogeneity estimator: Harrer et al. (2022), Doing Meta-Analysis with R, section 4.1.2.1, pp. 102–103; metafor rma.uni documentation (https://wviechtb.github.io/metafor/reference/rma.uni.html); verify and cite the original method literature.")
    if report["ci_method"] == "hksj":
        reminders.append("Modified Hartung–Knapp interval: Knapp and Hartung (2003), DOI 10.1002/sim.1482; verify the exact variant and its assumptions.")
    if result.get("prediction_interval_95"):
        reminders.append("The 95% prediction interval uses a k−2 degrees-of-freedom convention; verify and cite the convention used.")
    reminders.append("A software citation alone does not document the methods or the study evidence.")
    return reminders


def _heterogeneity_scope(result: dict) -> str:
    estimator = {"fixed": "not estimated for fixed effect", "random": "DerSimonian–Laird",
                 "random_pm": "Paule–Mandel", "random_reml": "REML"}[result["model"]]
    tau = ("τ² " + estimator + " = " + _number(result["tau2"]) +
           " on the " + result["effect_scale"] + " analysis scale") if result.get("tau2") is not None else "τ² " + estimator
    return ("I² uses fixed inverse-variance Q: max(0, (Q−df)/Q), with 0 when Q=0, "
            "regardless of pooling model. " + tau + ".")


def _report_title(report: dict) -> str:
    return f"Synthesis report: {report['outcome']}"


def _heterogeneity_interval_note(result: dict) -> str:
    interval = result.get("heterogeneity_intervals")
    if not interval:
        return ""
    if not interval["available"]:
        return f"Higgins–Thompson 95% intervals unavailable: {interval['reason']}"
    i2_low, i2_high = interval["i2_ci_percent"]
    h_low, h_high = interval["h_ci"]
    return (f"Higgins–Thompson 95% intervals (fixed inverse-variance Q): "
            f"I² {_number(i2_low)} to {_number(i2_high)}%; "
            f"H {_number(h_low)} to {_number(h_high)}. "
            f"Verify {interval['method_source']}.")


def build_docx(report: dict) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    result = report["result"]
    document = Document()
    document.sections[0].top_margin = Inches(0.65)
    document.sections[0].bottom_margin = Inches(0.65)
    document.add_heading(_report_title(report), 0)
    context = (f"{report['comparison']} · {report['outcome']} · {report['timepoint']} · "
               f"{report['measure']} · {result['n_studies']} studies")
    document.add_paragraph(context)
    document.add_heading("Pooled result", level=1)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.add_run(
        f"{report['measure']} {_number(result['pooled'])} "
        f"(95% CI {_number(result['ci_low'])} to {_number(result['ci_high'])})"
    ).bold = True
    document.add_paragraph(f"Model: {result['model']} — {result['method']}; confidence interval: {result['ci_method']}.")
    heterogeneity = (f"Q({result['q_df']}) = {_number(result['q'])}, p = {_number(result['q_p'])}; "
                     f"I² = {_number(result['i2_percent'])}%")
    if result.get("tau2") is not None:
        heterogeneity += f"; τ² = {_number(result['tau2'])}"
    document.add_paragraph(heterogeneity)
    document.add_paragraph(_heterogeneity_scope(result))
    if note := _heterogeneity_interval_note(result):
        document.add_paragraph(note)
    if result.get("prediction_interval_95"):
        low, high = result["prediction_interval_95"]
        document.add_paragraph(f"95% prediction interval: {_number(low)} to {_number(high)}")
    for warning in result.get("warnings", []):
        document.add_paragraph(f"Warning: {warning}")

    document.add_heading("Study estimates", level=1)
    document.add_paragraph("Study confidence intervals and the pooled interval are plotted on the analysis scale.")
    plot = document.add_table(rows=1, cols=3)
    plot.style = "Table Grid"
    plot.rows[0].cells[0].text = "Study"
    plot.rows[0].cells[1].text = "Estimate (95% CI)"
    plot.rows[0].cells[2].text = "Analysis-scale forest"
    bounds = _bounds(report)
    for effect in result["effects"]:
        center = _analysis_estimate(effect)
        low, high = center - 1.96 * effect["se"], center + 1.96 * effect["se"]
        shown_low = _display_value(low, effect["measure"])
        shown_high = _display_value(high, effect["measure"])
        cells = plot.add_row().cells
        cells[0].text = effect.get("study_label") or effect["study_id"]
        cells[1].text = f"{_number(_display_value(center, effect['measure']))} ({_number(shown_low)} to {_number(shown_high)})"
        cells[2].text = _plot_bar(low, center, high, bounds,
                                  reference=result["measure"] not in _POINT_ESTIMATES)
        for run in cells[2].paragraphs[0].runs:
            run.font.name = "Consolas"
            run.font.size = Pt(8)
    axis = document.add_paragraph(
        f"Analysis scale: {_number(bounds[0])} to {_number(bounds[1])}; "
        f"{'vertical bar = no effect (0); ' if result['measure'] not in _POINT_ESTIMATES else ''}"
        f"* = study estimate; line = normal 95% CI."
    )
    for run in axis.runs:
        run.font.size = Pt(8)

    document.add_heading("Selected study and source provenance", level=1)
    for effect in result["effects"]:
        document.add_heading(effect.get("study_label") or effect["study_id"], level=2)
        source = effect.get("source_title") or "Untitled report"
        authors_year = " ".join(str(value) for value in (effect.get("source_authors"), effect.get("source_year")) if value)
        if authors_year:
            source += f" ({authors_year})"
        details = [
            ("Study ID", effect["study_id"]),
            ("Selected result ID", effect["result_id"]),
            ("Source report", source),
            ("Source key", effect["source_key"]),
            ("Journal", effect.get("source_journal") or "Not recorded"),
            ("DOI", effect.get("source_doi") or ""),
            ("Source locator", effect.get("source_locator") or "Not recorded"),
            ("Entry method", effect.get("entry_method") or "Not recorded"),
            ("Saved effect / SE", f"{_number(effect['estimate'])} / {_number(effect['se'])} ({effect['measure']})"),
            ("Saved inputs", json.dumps(effect.get("input_data") or {}, ensure_ascii=False, sort_keys=True)),
        ]
        table = document.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        for key, value in details:
            cells = table.add_row().cells
            cells[0].text = key
            cells[1].text = str(value)

    document.add_heading("Citation reminders", level=1)
    for reminder in _citations(report):
        document.add_paragraph(reminder, style="List Bullet")
    out = BytesIO()
    document.save(out)
    return out.getvalue()


def build_pptx(report: dict) -> bytes:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt

    result = report["result"]
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]
    navy = RGBColor(31, 54, 76)
    blue = RGBColor(42, 111, 151)
    gray = RGBColor(106, 119, 130)

    def textbox(slide, x, y, w, h, text, size=16, bold=False, color=navy):
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = shape.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.margin_left = Inches(0.04)
        frame.margin_right = Inches(0.04)
        frame.margin_top = Inches(0.01)
        frame.margin_bottom = Inches(0.01)
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        paragraph = frame.paragraphs[0]
        paragraph.text = str(text)
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = color
        return shape

    def title(slide, text, subtitle=""):
        textbox(slide, .55, .3, 12.2, .55, text, 27, True)
        if subtitle:
            textbox(slide, .58, .88, 12.1, .42, subtitle, 12, False, gray)

    slide = presentation.slides.add_slide(blank)
    title(slide, _report_title(report),
          f"{report['comparison']} · {report['timepoint']} · {report['measure']} · {result['n_studies']} studies")
    textbox(slide, .7, 1.65, 5.8, .75,
            f"{report['measure']} {_number(result['pooled'])}  (95% CI {_number(result['ci_low'])} to {_number(result['ci_high'])})",
            22, True, blue)
    textbox(slide, .7, 2.55, 5.8, .7,
            f"{result['model']} · {result['method']}\n{result['ci_method']}", 16)
    heterogeneity = f"Q({result['q_df']}) = {_number(result['q'])}; p = {_number(result['q_p'])}; I² = {_number(result['i2_percent'])}%"
    if result.get("tau2") is not None:
        heterogeneity += f"; τ² = {_number(result['tau2'])}"
    textbox(slide, .7, 3.55, 5.8, .8, heterogeneity, 15)
    textbox(slide, 6.65, 3.48, 5.9, 1.1, _heterogeneity_scope(result), 11, False, gray)
    if note := _heterogeneity_interval_note(result):
        textbox(slide, 6.65, 4.45, 5.9, .55, note, 10, False, gray)
    if result.get("prediction_interval_95"):
        lo, hi = result["prediction_interval_95"]
        textbox(slide, .7, 4.42, 5.8, .48, f"95% prediction interval: {_number(lo)} to {_number(hi)}", 14)
    warnings = result.get("warnings", [])
    textbox(slide, .7, 5.05, 11.9, 1.05,
            "Warnings: " + ("; ".join(warnings) if warnings else "None returned by the synthesis method."),
            14, False, gray)
    textbox(slide, .7, 6.52, 11.9, .38, "Study labels, source reports, locators, and saved inputs follow on the next slides.", 11, False, gray)

    bounds = _bounds(report)
    effects = result["effects"]
    for page_start in range(0, len(effects), 8):
        page = presentation.slides.add_slide(blank)
        page_rows = effects[page_start:page_start + 8]
        reference_label = "vertical line marks no effect (0) · " if result["measure"] not in _POINT_ESTIMATES else ""
        title(page, "Forest plot", f"Analysis scale; right-hand values use the measure scale · {reference_label}normal 95% study limits · page {page_start // 8 + 1}")
        x0, x1 = 4.15, 10.6

        def x_pos(value: float) -> float:
            return x0 + (value - bounds[0]) / (bounds[1] - bounds[0]) * (x1 - x0)

        if result["measure"] not in _POINT_ESTIMATES:
            reference_x = x_pos(0)
            page.shapes.add_connector(1, Inches(reference_x), Inches(1.55), Inches(reference_x), Inches(6.5)).line.color.rgb = gray
        for index, effect in enumerate(page_rows):
            y = 1.85 + index * .55
            center = _analysis_estimate(effect)
            low, high = center - 1.96 * effect["se"], center + 1.96 * effect["se"]
            shown_low = _display_value(low, effect["measure"])
            shown_high = _display_value(high, effect["measure"])
            label = effect.get("study_label") or effect["study_id"]
            textbox(page, .58, y - .16, 3.4, .32, label[:48], 12)
            page.shapes.add_connector(1, Inches(x_pos(low)), Inches(y), Inches(x_pos(high)), Inches(y)).line.color.rgb = blue
            marker = page.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x_pos(center) - .075), Inches(y - .075), Inches(.15), Inches(.15))
            marker.fill.solid()
            marker.fill.fore_color.rgb = blue
            marker.line.color.rgb = blue
            textbox(page, 10.85, y - .16, 2.0, .32,
                    f"{_number(_display_value(center, effect['measure']))} ({_number(shown_low)} to {_number(shown_high)})", 10, False, navy)

        if page_start == 0:
            pooled_low, pooled_high = result["ci_analysis_low"], result["ci_analysis_high"]
            y = 6.05
            textbox(page, .58, y - .16, 3.4, .32, "Pooled", 12, True)
            page.shapes.add_connector(1, Inches(x_pos(pooled_low)), Inches(y), Inches(x_pos(pooled_high)), Inches(y)).line.color.rgb = navy
            diamond = page.shapes.add_shape(MSO_SHAPE.DIAMOND, Inches(x_pos(result['pooled_analysis']) - .09), Inches(y - .09), Inches(.18), Inches(.18))
            diamond.fill.solid()
            diamond.fill.fore_color.rgb = navy
            diamond.line.color.rgb = navy
            textbox(page, 10.85, y - .16, 2.0, .32, f"{_number(result['pooled'])} ({_number(result['ci_low'])} to {_number(result['ci_high'])})", 10, True, navy)
        textbox(page, x0, 6.52, x1 - x0, .3,
                f"{_number(bounds[0])}                                      {_number(bounds[1])}", 10, False, gray)
        textbox(page, .58, 6.82, 12.2, .24,
                f"{result['measure']} · {result['model']} · {result['ci_method']} · "
                f"τ²={_number(result.get('tau2'))} ({result['effect_scale']} scale)", 9, False, gray)
        textbox(page, .58, 7.05, 12.2, .37,
                "Selected result IDs: " + ", ".join(str(effect['result_id']) for effect in page_rows) +
                ". Study sources and saved inputs follow; verify and cite original methods and reports.", 9, False, gray)

    for effect in effects:
        page = presentation.slides.add_slide(blank)
        label = effect.get("study_label") or effect["study_id"]
        title(page, f"Study provenance: {label[:70]}", f"Study ID {effect['study_id']}")
        source_title = effect.get("source_title") or "Untitled report"
        authors_year = " ".join(str(value) for value in (effect.get("source_authors"), effect.get("source_year")) if value)
        lines = [
            ("Selected result ID", effect["result_id"]),
            ("Source report", f"{source_title} {f'({authors_year})' if authors_year else ''}"),
            ("Source key", effect["source_key"]),
            ("Journal", effect.get("source_journal") or "Not recorded"),
            ("DOI", effect.get("source_doi") or "Not recorded"),
            ("Source locator", effect.get("source_locator") or "Not recorded"),
            ("Entry method", effect.get("entry_method") or "Not recorded"),
            ("Saved effect / SE", f"{_number(effect['estimate'])} / {_number(effect['se'])} ({effect['measure']})"),
        ]
        y = 1.35
        for key, value in lines:
            textbox(page, .7, y, 2.2, .42, key, 13, True, gray)
            textbox(page, 2.9, y, 9.7, .54, value, 16)
            y += .62
        textbox(page, .7, 6.48, 11.8, .38, "Complete saved inputs follow on the next slide(s).", 11, False, gray)

        saved = json.dumps(effect.get("input_data") or {}, ensure_ascii=False, sort_keys=True, indent=2)
        input_lines = [part for line in saved.splitlines() for part in
                       (textwrap.wrap(line, width=70, replace_whitespace=False,
                                      drop_whitespace=False, break_on_hyphens=False) or [""])]
        for start in range(0, len(input_lines), 25):
            input_page = presentation.slides.add_slide(blank)
            title(input_page, f"Saved inputs: {label[:70]}",
                  f"Result {effect['result_id']} · source {effect['source_key']} · page {start // 25 + 1}")
            shape = textbox(input_page, .8, 1.4, 11.8, 5.7,
                            "\n".join(input_lines[start:start + 25]), 11)
            shape.text_frame.vertical_anchor = MSO_ANCHOR.TOP
            shape.text_frame.paragraphs[0].font.name = "Consolas"

    reminders = _citations(report)
    for start in range(0, len(reminders), 3):
        page = presentation.slides.add_slide(blank)
        title(page, "Citation reminders", f"Page {start // 3 + 1}")
        for index, reminder in enumerate(reminders[start:start + 3]):
            textbox(page, .8, 1.45 + index * 1.7, 11.8, 1.5, f"• {reminder}", 13)
    out = BytesIO()
    presentation.save(out)
    return out.getvalue()
